"""
Gerador de Pareceres Atuariais — Backend
Lumens Atuarial | Núcleo Judicial
"""

from flask import Flask, jsonify, request, send_file, send_from_directory, session
from flask_cors import CORS
import dropbox
from dropbox.exceptions import AuthError, ApiError
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy
import io
import os
import functools

# ── ENV ──────────────────────────────────────────────────────────────────────

def carregar_env():
    env_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), '.env')
    if os.path.exists(env_path):
        with open(env_path, 'r', encoding='utf-8') as f:
            for linha in f:
                linha = linha.strip()
                if linha and '=' in linha and not linha.startswith('#'):
                    chave, valor = linha.split('=', 1)
                    os.environ[chave.strip()] = valor.strip()

carregar_env()

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "lumens-judicial-chave-2026")
CORS(app, supports_credentials=True, origins=["https://criador-de-pareceres.onrender.com", "http://localhost:5000"])

# ── HELPERS ───────────────────────────────────────────────────────────────────

def get_dropbox_client():
    refresh_token = os.environ.get("DROPBOX_REFRESH_TOKEN", "").strip()
    app_key       = os.environ.get("DROPBOX_APP_KEY", "").strip()
    app_secret    = os.environ.get("DROPBOX_APP_SECRET", "").strip()
    if refresh_token and app_key and app_secret:
        return dropbox.Dropbox(
            oauth2_refresh_token=refresh_token,
            app_key=app_key,
            app_secret=app_secret
        )
    token = os.environ.get("DROPBOX_TOKEN", "").strip()
    if token:
        return dropbox.Dropbox(token)
    raise ValueError("Credenciais do Dropbox não configuradas.")

def get_pasta():
    return os.environ.get("DROPBOX_PASTA", "/Banco de Teses").strip()

def login_required(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("autenticado"):
            return jsonify({"erro": "Não autenticado"}), 401
        return f(*args, **kwargs)
    return decorated

# ── AUTENTICAÇÃO ──────────────────────────────────────────────────────────────

@app.route("/api/check", methods=["GET"])
def check():
    return jsonify({"autenticado": bool(session.get("autenticado"))})

@app.route("/api/login", methods=["POST"])
def rota_login():
    data  = request.get_json() or {}
    senha = data.get("senha", "")
    senha_correta = os.environ.get("APP_SENHA", "JudicialLumens01")
    if senha == senha_correta:
        session["autenticado"] = True
        session.permanent = False
        return jsonify({"ok": True})
    return jsonify({"erro": "Senha incorreta"}), 401

@app.route("/api/logout", methods=["POST"])
def rota_logout():
    session.clear()
    return jsonify({"ok": True})

# ── FRONTEND ──────────────────────────────────────────────────────────────────

@app.route("/")
def index():
    pasta = os.path.dirname(os.path.abspath(__file__))
    return send_from_directory(pasta, 'criador-pareceres.html')

# ── STATUS ────────────────────────────────────────────────────────────────────

@app.route("/api/status", methods=["GET"])
def status():
    template_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'template.docx')
    dropbox_ok = bool(
        os.environ.get("DROPBOX_REFRESH_TOKEN") or os.environ.get("DROPBOX_TOKEN")
    )
    return jsonify({
        "status": "ok",
        "dropbox_configurado": dropbox_ok,
        "template_ok": os.path.exists(template_path),
        "pasta": get_pasta()
    })

# ── DIAGNÓSTICO ───────────────────────────────────────────────────────────────

@app.route("/api/explorar", methods=["GET"])
def explorar():
    try:
        dbx  = get_dropbox_client()
        path = request.args.get("path", "")
        resultado = dbx.files_list_folder(path, recursive=False)
        itens = [{"nome": e.name, "path": e.path_display} for e in resultado.entries]
        return jsonify({"path_consultado": path if path else "/", "itens": itens})
    except Exception as e:
        return jsonify({"erro": str(e)}), 500

# ── TÓPICOS ───────────────────────────────────────────────────────────────────

# Pastas que têm prioridade fixa no topo da lista (case-insensitive)
PASTAS_PRIORIDADE = ["base", "quesitos"]

def _sort_key_categoria(cat_lower):
    """
    Retorna uma tupla de ordenação:
      (0, índice) para pastas prioritárias
      (1, nome)   para as demais em ordem alfabética
    """
    try:
        idx = PASTAS_PRIORIDADE.index(cat_lower)
        return (0, idx, "")
    except ValueError:
        return (1, 999, cat_lower)

@app.route("/api/topicos", methods=["GET"])
@login_required
def listar_topicos():
    try:
        dbx   = get_dropbox_client()
        pasta = get_pasta()

        # Buscar com paginação completa (evita perder itens quando has_more=True)
        resultado = dbx.files_list_folder(pasta, recursive=True)
        todas_entradas = list(resultado.entries)
        while resultado.has_more:
            resultado = dbx.files_list_folder_continue(resultado.cursor)
            todas_entradas.extend(resultado.entries)

        topicos = []
        idx = 1
        pasta_lower = pasta.lower().rstrip("/")

        for entry in todas_entradas:
            if isinstance(entry, dropbox.files.FileMetadata):
                if entry.name.lower().endswith(".docx"):
                    # Caminho relativo usando path_display (preserva acentos e maiúsculas)
                    pd = entry.path_display
                    prefixo = pasta.rstrip("/") + "/"
                    if pd.lower().startswith(prefixo.lower()):
                        rel_display = pd[len(prefixo):]
                    else:
                        rel_display = pd
                    partes_display = rel_display.split("/")

                    if len(partes_display) > 1:
                        categoria_display = partes_display[0]
                        categoria_lower   = partes_display[0].lower()
                    else:
                        categoria_display = "Geral"
                        categoria_lower   = "geral"

                    nome = entry.name
                    for ext in (".docx", ".DOCX", ".Docx"):
                        nome = nome.replace(ext, "")

                    # Flags especiais
                    eh_topico_principal = (
                        categoria_lower == "base" and
                        nome.lower() == "cálculo lumens"
                    )
                    eh_ultima_pagina = nome.lower() in ("anexo", "apêndice", "apendice")

                    topicos.append({
                        "id":              idx,
                        "nome":            nome,
                        "categoria":       categoria_display,
                        "categoria_lower": categoria_lower,
                        "path":            entry.path_display,
                        "topico_principal": eh_topico_principal,
                        "ultima_pagina":   eh_ultima_pagina,
                    })
                    idx += 1

        # Ordenação: prioridade de pasta → alfabética de nome
        topicos.sort(key=lambda x: (
            _sort_key_categoria(x["categoria_lower"]),
            x["nome"].lower()
        ))

        return jsonify({"topicos": topicos})

    except AuthError:
        return jsonify({"erro": "Token do Dropbox inválido ou expirado."}), 401
    except ApiError as e:
        return jsonify({"erro": f"Erro na API do Dropbox: {str(e)}"}), 500
    except ValueError as e:
        return jsonify({"erro": str(e)}), 400
    except Exception as e:
        return jsonify({"erro": f"Erro inesperado: {str(e)}"}), 500

# ── GERAÇÃO DO PARECER ────────────────────────────────────────────────────────

@app.route("/api/gerar", methods=["POST"])
@login_required
def gerar_parecer():
    data                = request.get_json() or {}
    topicos_selecionados = data.get("topicos", [])
    dados_caso          = data.get("dados", {})

    if not topicos_selecionados:
        return jsonify({"erro": "Nenhum tópico selecionado."}), 400

    try:
        dbx           = get_dropbox_client()
        template_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'template.docx')

        if not os.path.exists(template_path):
            return jsonify({"erro": "template.docx não encontrado no servidor."}), 500

        doc_final = Document(template_path)

        # Substituir placeholders da capa no template
        _substituir_dados(doc_final.element.body, dados_caso)

        # Separar os tópicos em dois grupos:
        #   1. normais    → entram antes do bloco de encerramento (assinatura)
        #   2. ultima_pag → "Anexo" / "Apêndice": entram após a quebra de página do encerramento
        normais    = [t for t in topicos_selecionados if not t.get("ultima_pagina")]
        ultima_pag = [t for t in topicos_selecionados if t.get("ultima_pagina")]

        # Ponto de inserção dos tópicos normais:
        # sempre ANTES do bloco "É o Parecer Técnico." / assinatura
        ponto = _encontrar_ponto_insercao(doc_final)

        for topico in normais:
            docx_topico = _baixar_docx(dbx, topico["path"])
            _copiar_estilos(docx_topico, doc_final)
            eh_principal = topico.get("topico_principal", False)
            _inserir_topico(docx_topico, doc_final, dados_caso, ponto, eh_principal)
            # Atualizar o ponto para inserir o próximo tópico após o anterior
            ponto = None  # após o primeiro, inserir em sequência antes do encerramento

        # Tópicos de última página: inseridos APÓS a quebra de página do encerramento
        for topico in ultima_pag:
            docx_topico = _baixar_docx(dbx, topico["path"])
            _copiar_estilos(docx_topico, doc_final)
            _inserir_ultima_pagina(docx_topico, doc_final, dados_caso)

        buffer = io.BytesIO()
        doc_final.save(buffer)
        buffer.seek(0)

        processo     = dados_caso.get("processo", "sp").strip()
        # Limpar caracteres inválidos para nome de arquivo, preservando o número do processo
        import re as _re
        processo_limpo = _re.sub(r'[\\/*?:"<>|]', '-', processo)
        nome_arquivo = f"{processo_limpo}_Parecer Técnico.rev001.docx"

        return send_file(
            buffer,
            as_attachment=True,
            download_name=nome_arquivo,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    except AuthError:
        return jsonify({"erro": "Token do Dropbox inválido ou expirado."}), 401
    except Exception as e:
        import traceback
        return jsonify({"erro": f"Erro ao gerar parecer: {str(e)}", "detalhe": traceback.format_exc()}), 500

# ── FUNÇÕES AUXILIARES ────────────────────────────────────────────────────────

def _baixar_docx(dbx, path):
    _, response = dbx.files_download(path)
    return Document(io.BytesIO(response.content))


def _copiar_estilos(doc_origem, doc_destino):
    estilos_origem  = doc_origem.element.find(qn("w:styles"))
    estilos_destino = doc_destino.element.find(qn("w:styles"))
    if estilos_origem is None or estilos_destino is None:
        return
    ids_existentes = {
        e.get(qn("w:styleId"))
        for e in estilos_destino.findall(qn("w:style"))
        if e.get(qn("w:styleId"))
    }
    for estilo in estilos_origem.findall(qn("w:style")):
        sid = estilo.get(qn("w:styleId"))
        if sid and sid not in ids_existentes:
            estilos_destino.append(copy.deepcopy(estilo))


def _encontrar_ponto_insercao(doc):
    """
    Retorna o elemento que serve como ponto de inserção dos tópicos.
    Procura o placeholder [inserir primeira impugnação].
    Se não achar, usa o elemento imediatamente ANTES do bloco de encerramento
    ("É o Parecer Técnico."), garantindo que os tópicos entrem antes da assinatura.
    """
    body = doc.element.body

    # Prioridade 1: placeholder explícito
    for elem in body:
        for t in elem.iter(qn("w:t")):
            if t.text and "[inserir primeira impugna" in t.text:
                return elem

    # Prioridade 2: elemento imediatamente antes de "É o Parecer Técnico."
    children = list(body)
    for i, elem in enumerate(children):
        for t in elem.iter(qn("w:t")):
            if t.text and "É o Parecer Técnico" in t.text:
                # Retornar o elemento anterior como ponto de inserção
                return children[i - 1] if i > 0 else elem

    # Fallback: antes do sectPr
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is not None:
        idx = children.index(sect_pr)
        return children[idx - 1] if idx > 0 else None
    return None


def _encontrar_inicio_encerramento(doc):
    """
    Retorna o índice do primeiro elemento do bloco de encerramento
    (parágrafo vazio antes de "É o Parecer Técnico.", ou o próprio).
    Usado para inserir tópicos ANTES desse bloco.
    """
    body     = doc.element.body
    children = list(body)
    for i, elem in enumerate(children):
        for t in elem.iter(qn("w:t")):
            if t.text and "É o Parecer Técnico" in t.text:
                # Incluir parágrafo vazio imediatamente anterior se existir
                if i > 0:
                    prev_texts = [t.text for t in children[i-1].iter(qn("w:t")) if t.text]
                    if not prev_texts:
                        return i - 1
                return i
    return None


def _encontrar_pos_encerramento(doc):
    """
    Retorna o índice do elemento APÓS o bloco completo de encerramento
    (assinatura + quebra de página), que é onde Anexo/Apêndice do banco devem entrar.
    O bloco de encerramento termina no elemento com pageBreak.
    """
    body     = doc.element.body
    children = list(body)
    # Achar o pageBreak que vem após a assinatura
    em_encerramento = False
    for i, elem in enumerate(children):
        # Detectar início do encerramento
        for t in elem.iter(qn("w:t")):
            if t.text and "É o Parecer Técnico" in t.text:
                em_encerramento = True
        # Detectar quebra de página após o encerramento
        if em_encerramento:
            tem_page_break = any(
                br.get(qn("w:type")) == "page"
                for br in elem.iter(qn("w:br"))
            )
            if tem_page_break:
                return i + 1  # posição APÓS a quebra
    return None


def _inserir_topico(doc_topico, doc_final, dados_caso, ponto_ref, eh_principal):
    """
    Insere o conteúdo de doc_topico ANTES do bloco de encerramento
    ("É o Parecer Técnico." + assinatura).
    Se ponto_ref for o placeholder [inserir primeira impugnação], substitui-o.
    Se ponto_ref for None, localiza automaticamente o início do encerramento.
    """
    body_origem  = doc_topico.element.body
    body_destino = doc_final.element.body

    elementos = [e for e in body_origem if e.tag != qn("w:sectPr")]
    if not elementos:
        return

    children = list(body_destino)

    if ponto_ref is not None and ponto_ref in children:
        pos = children.index(ponto_ref)
        placeholder_encontrado = any(
            "[inserir primeira impugna" in (t.text or "")
            for t in ponto_ref.iter(qn("w:t"))
        )
        if placeholder_encontrado:
            insert_pos = pos
            body_destino.remove(ponto_ref)
        else:
            insert_pos = pos + 1
    else:
        # Inserir imediatamente antes do bloco de encerramento
        idx_enc = _encontrar_inicio_encerramento(doc_final)
        if idx_enc is not None:
            insert_pos = idx_enc
        else:
            # fallback: antes do sectPr
            sect_pr = body_destino.find(qn("w:sectPr"))
            insert_pos = list(body_destino).index(sect_pr) if sect_pr is not None else len(list(body_destino))

    for i, elem in enumerate(elementos):
        novo = copy.deepcopy(elem)
        _substituir_dados(novo, dados_caso)
        if i == 0:
            _ajustar_estilo_titulo(novo, eh_principal)
        body_destino.insert(insert_pos + i, novo)


def _ajustar_estilo_titulo(elem, eh_principal):
    """
    Se o primeiro parágrafo do tópico tiver um estilo de título,
    promove para cTTULONVEL1 (tópico) ou mantém dTTULONVEL2 (subtópico).
    """
    pPr   = elem.find(qn("w:pPr"))
    if pPr is None:
        return
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is None:
        return
    estilo_atual = pStyle.get(qn("w:val"), "")
    # Se já é um estilo de título do template, ajustar conforme tipo
    estilos_titulo = {"cTTULONVEL1", "dTTULONVEL2", "Heading1", "Heading2"}
    if estilo_atual in estilos_titulo or "TTULO" in estilo_atual.upper() or "TITULO" in estilo_atual.upper():
        novo_estilo = "cTTULONVEL1" if eh_principal else "dTTULONVEL2"
        pStyle.set(qn("w:val"), novo_estilo)


def _inserir_ultima_pagina(doc_topico, doc_final, dados_caso):
    """
    Insere Anexo/Apêndice APÓS a quebra de página do bloco de encerramento
    (assinatura), sempre ao final do documento — antes do sectPr.
    Preserva os estilos originais do arquivo sem alteração.
    """
    body_origem  = doc_topico.element.body
    body_destino = doc_final.element.body
    sect_pr      = body_destino.find(qn("w:sectPr"))

    # Posição de inserção: após o pageBreak do encerramento, antes do sectPr
    pos_enc = _encontrar_pos_encerramento(doc_final)
    if pos_enc is None:
        # fallback: antes do sectPr
        children = list(body_destino)
        pos_enc = children.index(sect_pr) if sect_pr is not None else len(children)

    def inserir_em(elem, offset):
        body_destino.insert(pos_enc + offset, elem)

    offset = 0
    for elem in body_origem:
        if elem.tag == qn("w:sectPr"):
            continue
        novo = copy.deepcopy(elem)
        _substituir_dados(novo, dados_caso)
        inserir_em(novo, offset)
        offset += 1


def _substituir_dados(elemento, dados_caso):
    """Substitui placeholders no XML do elemento."""
    mapeamento = {
        "[demanda]":         dados_caso.get("demanda", ""),
        "[nº do processo]":  dados_caso.get("processo", ""),
        "[Autor/Reclamante]": dados_caso.get("participante", ""),
        "[Vara/Juízo]":      dados_caso.get("vara", ""),
        "[data de entrega]": dados_caso.get("entrega", ""),
        # aliases mais antigos (compatibilidade)
        "{{participante}}":  dados_caso.get("participante", ""),
        "{{processo}}":      dados_caso.get("processo", ""),
        "{{vara}}":          dados_caso.get("vara", ""),
        "{{demanda}}":       dados_caso.get("demanda", ""),
        "{{entrega}}":       dados_caso.get("entrega", ""),
    }
    for no_texto in elemento.iter(qn("w:t")):
        if no_texto.text:
            for placeholder, valor in mapeamento.items():
                if placeholder in no_texto.text and valor:
                    no_texto.text = no_texto.text.replace(placeholder, valor)


# ── INICIALIZAÇÃO ─────────────────────────────────────────────────────────────

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
